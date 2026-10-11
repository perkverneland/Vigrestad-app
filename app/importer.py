"""Import og vasking av gamle data fra Prosess.xlsm.

Alle registreringer spilles av i datorekkefølge med de samme reglene som appen bruker
(hele lotet flyttes videre til nytt steg). Avvik skrives til en rapport.
"""
from collections import defaultdict
from datetime import date, datetime

from . import lager
from .lager import Feil, normaliser


def _til_dato(verdi):
    if isinstance(verdi, datetime):
        return verdi.date().isoformat()
    if isinstance(verdi, date):
        return verdi.isoformat()
    return None


def _til_tall(verdi):
    if isinstance(verdi, (int, float)):
        return float(verdi)
    try:
        return float(str(verdi).strip().replace(",,", ",").replace(",", "."))
    except (TypeError, ValueError):
        return None


def _lotnr(verdi):
    if verdi is None:
        return None
    if isinstance(verdi, float) and verdi.is_integer():
        verdi = int(verdi)
    s = str(verdi).strip()
    return s or None


class Oppslag:
    """Finner stamdata-ID ut fra fritekst, uavhengig av store/små bokstaver og mellomrom."""

    def __init__(self, db, tabell):
        self.db, self.tabell = db, tabell
        self.kart = {normaliser(r["navn"]): r["id"] for r in db.execute(f"SELECT id, navn FROM {tabell}")}

    def finn(self, navn, opprett=True, **ekstra):
        nokkel = normaliser(navn)
        if not nokkel:
            return None
        if nokkel not in self.kart and opprett:
            pent = " ".join(str(navn).split())
            pent = pent[0].upper() + pent[1:]
            kolonner = ["navn", *ekstra]
            self.kart[nokkel] = self.db.execute(
                f"INSERT INTO {self.tabell} ({', '.join(kolonner)}) VALUES ({', '.join('?' * len(kolonner))})",
                (pent, *ekstra.values())).lastrowid
        return self.kart.get(nokkel)


def importer(db, sti, logg=print):
    import openpyxl

    logg(f"Leser {sti} ...")
    wb = openpyxl.load_workbook(sti, data_only=True, read_only=True)
    rapport = ["IMPORTRAPPORT – avvik funnet i Prosess.xlsm", ""]
    avvik = defaultdict(list)

    def ark(navn):
        for kandidat in wb.sheetnames:
            if kandidat.strip().lower() == navn.lower():
                return list(wb[kandidat].iter_rows(min_row=2, values_only=True))
        raise Feil(f"Fant ikke arket «{navn}» i Excel-filen.")

    def er_overskrift(rad):
        return isinstance(rad[0], str) and rad[0].strip().lower() == "dato"

    spes = Oppslag(db, "spesifikasjoner")
    prod = Oppslag(db, "produkter")

    # Stamdata fra «Lister»: produkter med varenr
    for rad in ark("Lister"):
        if len(rad) >= 5 and rad[3] and rad[4] and normaliser(rad[3]) not in prod.kart:
            prod.finn(rad[3], varenr=str(rad[4]))

    # ---- Lot
    lots = {}
    for nr, rad in enumerate(ark("Database Lot"), start=2):
        if er_overskrift(rad) or not any(rad[:5]):
            continue
        lotnr = _lotnr(rad[2])
        if not lotnr:
            avvik["Lot uten lotnummer (hoppet over)"].append(f"rad {nr}: {rad[:5]}")
            continue
        if lotnr in lots:
            avvik["Lotnummer registrert flere ganger (første beholdt)"].append(f"{lotnr} (rad {nr})")
            continue
        if len(lotnr) != 10 or not lotnr.isdigit():
            avvik["Lotnummer med uvanlig format"].append(lotnr)
        lots[lotnr] = lager.opprett_lot(
            db, lotnr=lotnr, dato=_til_dato(rad[0]), prima_lotnr=_lotnr(rad[1]),
            spesifikasjon_id=spes.finn(rad[3]), produkt_id=prod.finn(rad[4]),
            kilde="import", krev_felt=False)
    logg(f"  {len(lots)} lot")

    def lot_id(lotnr, dato):
        if lotnr not in lots:
            lots[lotnr] = lager.opprett_lot(db, lotnr=lotnr, dato=dato, mangler_mottak=True,
                                            kilde="import", krev_felt=False)
            avvik["Lot som bare finnes i operasjons-/korrigeringsloggen (opprettet automatisk)"].append(lotnr)
        return lots[lotnr]

    # ---- Operasjoner
    operasjoner = []
    ukjente_op = defaultdict(int)
    for nr, rad in enumerate(ark("Database Operasjon"), start=2):
        if er_overskrift(rad) or not any(rad[:4]):
            continue
        dato, lotnr, kg = _til_dato(rad[0]), _lotnr(rad[1]), _til_tall(rad[3])
        op = lager.finn_operasjon(db, rad[2]) if rad[2] else None
        if not dato or not lotnr:
            avvik["Operasjon uten dato eller lotnr (hoppet over)"].append(f"rad {nr}: {rad[:4]}")
            continue
        if op is None:
            ukjente_op[str(rad[2])] += 1
            continue
        if kg is None or kg < 0:
            avvik["Operasjon med ugyldig kg (hoppet over)"].append(f"rad {nr}: {lotnr} {rad[2]} kg={rad[3]!r}")
            continue
        if kg == 0:
            avvik["Operasjon med 0 kg (importert)"].append(f"rad {nr}: {lotnr} {op['navn']}")
        operasjoner.append({"nr": nr, "dato": dato, "lotnr": lotnr, "op": op, "kg": kg, "brukt": False})
    for navn, antall in ukjente_op.items():
        avvik["Ukjent operasjon (hoppet over)"].append(f"«{navn}» ({antall} rader)")

    # ---- Korrigering/omlotting: minuspost på fra-lot fulgt av plusspost på til-lot
    korr = [(nr, _til_dato(r[0]), _lotnr(r[1]), _til_tall(r[2]))
            for nr, r in enumerate(ark("Database Korrigering"), start=2)
            if not er_overskrift(r) and r[1] is not None]
    omlottinger = []
    i = 0
    while i < len(korr):
        nr, dato, lotnr, kg = korr[i]
        neste = korr[i + 1] if i + 1 < len(korr) else None
        if kg is not None and kg < 0 and neste and neste[3] is not None and abs(neste[3] + kg) < 0.01:
            omlottinger.append({"nr": nr, "dato": dato, "fra": lotnr, "til": neste[2], "kg": -kg})
            i += 2
        else:
            avvik["Korrigering uten motpost (hoppet over)"].append(f"rad {nr}: {lotnr} {kg}")
            i += 1
    # Finn operasjonen som ble registrert på til-lotet samtidig (makroen la den i operasjonsloggen)
    etter_nokkel = defaultdict(list)
    for o in operasjoner:
        etter_nokkel[(o["dato"], o["lotnr"])].append(o)
    for om in omlottinger:
        for o in etter_nokkel[(om["dato"], om["til"])]:
            if not o["brukt"] and abs(o["kg"] - om["kg"]) < 0.01:
                o["brukt"] = True
                om["op"] = o
                break

    # ---- Svinn
    svinn = []
    for nr, rad in enumerate(ark("Database Svinn"), start=2):
        if er_overskrift(rad) or not rad[1]:
            continue
        svinn.append({"nr": nr, "dato": _til_dato(rad[0]), "lotnr": _lotnr(rad[1]), "kg": _til_tall(rad[3])})

    # ---- Spill av alt i datorekkefølge
    hendelser = ([(o["dato"], 0, o["nr"], "op", o) for o in operasjoner if not o["brukt"]]
                 + [(om["dato"], 1, om["nr"], "om", om) for om in omlottinger]
                 + [(s["dato"], 2, s["nr"], "svinn", s) for s in svinn])
    hendelser.sort(key=lambda h: (h[0], h[1], h[2]))
    label = lager.finn_operasjon(db, "Label kg")
    for dato, _, nr, type_, h in hendelser:
        try:
            if type_ == "op":
                lid = lot_id(h["lotnr"], dato)
                if h["op"]["id"] == label["id"]:
                    db.execute("UPDATE lots SET start_kg = COALESCE(start_kg, ?) WHERE id = ?", (h["kg"], lid))
                lager.registrer_operasjon(db, lot_id=lid, operasjon_id=h["op"]["id"], dato=dato, kg=h["kg"],
                                          kilde="import", tillat_null=True)
            elif type_ == "om":
                fra, til = lot_id(h["fra"], dato), lot_id(h["til"], dato)
                apne = lager.apne_posisjoner(db, fra)
                op = h.get("op")
                if not apne:
                    avvik["Omlotting fra lot uten beholdning (registrert uten å trekke fra)"].append(
                        f"rad {nr}: {h['fra']} -> {h['til']} {h['kg']} kg")
                kg_inn = op["kg"] if op else h["kg"]
                op_id = op["op"]["id"] if op else (apne[0]["operasjon_id"] if apne else label["id"])
                if not apne:
                    # Ingen beholdning på fra-lotet: legg kg inn på til-lotet og noter koblingen
                    hid = lager.registrer_operasjon(db, lot_id=til, operasjon_id=op_id, dato=dato, kg=kg_inn,
                                                    uten_kilde=True, kilde="import", tillat_null=True)
                    db.execute("UPDATE hendelser SET type = 'omlotting', fra_lot_id = ? WHERE id = ?", (fra, hid))
                    continue
                for_kg = apne[0]["kg"]
                delvis = h["kg"] < for_kg - 0.01
                lager.registrer_operasjon(db, lot_id=til, fra_lot_id=fra, operasjon_id=op_id, dato=dato,
                                          kg=kg_inn, delvis=delvis, rest_kg=round(for_kg - h["kg"], 3) if delvis else None,
                                          kilde="import", tillat_null=True)
            else:
                lid = lot_id(h["lotnr"], dato)
                lager.registrer_svinn(db, lot_id=lid, kg=h["kg"], dato=dato, kilde="import")
        except Feil as e:
            avvik["Registrering som ikke kunne importeres"].append(f"{type_} rad {nr}: {e}")

    db.commit()
    logg(f"  {len(operasjoner)} operasjoner, {len(omlottinger)} omlottinger, {len(svinn)} svinn")
    for overskrift, linjer in avvik.items():
        rapport.append(f"== {overskrift}: {len(linjer)}")
        rapport.extend("   " + linje for linje in linjer)
        rapport.append("")
    if not avvik:
        rapport.append("Ingen avvik funnet.")
    return rapport
